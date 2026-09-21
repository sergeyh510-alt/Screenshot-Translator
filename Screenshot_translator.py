import base64
import io
import json
import os
import subprocess
import sys
import threading
import tkinter as tk
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import requests
from PIL import ImageGrab
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

APP_DIR = Path.home() / ".screenshot_translator"
SETTINGS_PATH = APP_DIR / "settings.xml"
LEGACY_SETTINGS_PATH = APP_DIR / "settings.json"
APP_DIR_PATH = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
COLORS = {
    "bg": "#f4f7fb",
    "surface": "#ffffff",
    "surface_alt": "#eef4fb",
    "primary": "#2563eb",
    "primary_dark": "#1d4ed8",
    "text": "#172033",
    "muted": "#64748b",
    "border": "#d7e0ec",
}
DEFAULT_SETTINGS = {
    "api_url": "https://routerai.net/api/v1/chat/completions",
    "api_key": "",
    "model": "",
    "prompt": "Recognize the text in the image. Return it strictly in this format:\nEN: <original sentence in English>\nRU: <exact Russian translation>.\nDo not add explanations.",
    "output_docx": str(Path.home() / "ScreenshotTranslations.docx"),
    "mode": "append",
    "margin_top": "2.0",
    "margin_bottom": "2.0",
    "margin_left": "2.0",
    "margin_right": "2.0",
    "ask_open_document": True,
}


def load_settings():
    APP_DIR.mkdir(parents=True, exist_ok=True)
    data = DEFAULT_SETTINGS.copy()
    if SETTINGS_PATH.exists():
        try:
            root = ET.parse(SETTINGS_PATH).getroot()
            for item in root:
                if item.tag in data and item.text is not None:
                    value = item.text
                    data[item.tag] = value.lower() == "true" if item.tag == "ask_open_document" else value
        except Exception:
            pass
    elif LEGACY_SETTINGS_PATH.exists():
        # Read settings created by older versions and migrate them to XML.
        try:
            data.update(json.loads(LEGACY_SETTINGS_PATH.read_text(encoding="utf-8")))
            save_settings(data)
        except Exception:
            pass
    return data


def save_settings(data):
    APP_DIR.mkdir(parents=True, exist_ok=True)
    root = ET.Element("screenshot_translator_settings", {"version": "1.0"})
    for key in DEFAULT_SETTINGS:
        item = ET.SubElement(root, key)
        item.text = str(data.get(key, DEFAULT_SETTINGS[key])).lower() if isinstance(data.get(key), bool) else str(data.get(key, DEFAULT_SETTINGS[key]))
    tree = ET.ElementTree(root)
    ET.indent(tree, space="  ")
    tree.write(SETTINGS_PATH, encoding="utf-8", xml_declaration=True)


class RegionSelector(tk.Toplevel):
    def __init__(self, master, callback):
        super().__init__(master)
        self.callback = callback
        self.start = None
        self.geometry(f"{self.winfo_screenwidth()}x{self.winfo_screenheight()}+0+0")
        self.overrideredirect(True)
        self.attributes("-alpha", 0.28)
        self.attributes("-topmost", True)
        self.canvas = tk.Canvas(self, bg="black", cursor="crosshair", highlightthickness=0)
        self.canvas.pack(fill="both", expand=True)
        self.canvas.create_text(20, 20, anchor="nw", fill="white", font=("Arial", 16, "bold"),
                                text="Drag a rectangle around the text with the mouse. Esc — cancel")
        self.canvas.bind("<ButtonPress-1>", self.begin)
        self.canvas.bind("<B1-Motion>", self.drag)
        self.canvas.bind("<ButtonRelease-1>", self.end)
        self.bind("<Escape>", lambda e: self.destroy())
        self.focus_force()

    def begin(self, event):
        self.start = (event.x, event.y)
        self.rect = self.canvas.create_rectangle(event.x, event.y, event.x, event.y,
                                                 outline="#36d399", width=3)

    def drag(self, event):
        if self.start:
            self.canvas.coords(self.rect, self.start[0], self.start[1], event.x, event.y)

    def end(self, event):
        if not self.start:
            return
        x1, y1 = self.start
        x2, y2 = event.x, event.y
        left, top, right, bottom = min(x1, x2), min(y1, y2), max(x1, x2), max(y1, y2)
        if right - left >= 5 and bottom - top >= 5:
            self.callback((left, top, right, bottom))
        self.destroy()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Screenshot Translator — AI")
        self.geometry("790x650")
        self.minsize(700, 560)
        self.setup_theme()
        self.set_app_icon()
        self.settings = load_settings()
        self.region = None
        self.busy = False
        self.api_url_var = tk.StringVar(value=self.settings.get("api_url", ""))
        self.api_key_var = tk.StringVar(value=self.settings.get("api_key", ""))
        self.model_var = tk.StringVar(value=self.settings.get("model", ""))
        self.margin_top_var = tk.StringVar(value=self.settings.get("margin_top", "2.0"))
        self.margin_bottom_var = tk.StringVar(value=self.settings.get("margin_bottom", "2.0"))
        self.margin_left_var = tk.StringVar(value=self.settings.get("margin_left", "2.0"))
        self.margin_right_var = tk.StringVar(value=self.settings.get("margin_right", "2.0"))
        self.ask_open_document = tk.BooleanVar(value=self.settings.get("ask_open_document", True))
        self.build_ui()

    def setup_theme(self):
        self.configure(bg=COLORS["bg"])
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("TFrame", background=COLORS["bg"])
        style.configure("TLabel", background=COLORS["bg"], foreground=COLORS["text"], font=("Segoe UI", 10))
        style.configure("Header.TLabel", background=COLORS["bg"], foreground=COLORS["primary_dark"], font=("Segoe UI", 19, "bold"))
        style.configure("Subtitle.TLabel", background=COLORS["bg"], foreground=COLORS["muted"], font=("Segoe UI", 10))
        style.configure("TButton", padding=(12, 7), font=("Segoe UI", 10))
        style.configure("Accent.TButton", background=COLORS["primary"], foreground="white", borderwidth=0, padding=(14, 9), font=("Segoe UI", 10, "bold"))
        style.map("Accent.TButton", background=[("active", COLORS["primary_dark"]), ("pressed", COLORS["primary_dark"])])
        style.configure("TNotebook", background=COLORS["bg"], borderwidth=0)
        style.configure("TNotebook.Tab", background=COLORS["surface_alt"], foreground=COLORS["text"], padding=(16, 8))
        style.map("TNotebook.Tab", background=[("selected", COLORS["surface"])], foreground=[("selected", COLORS["primary_dark"])])
        style.configure("TEntry", fieldbackground=COLORS["surface"], foreground=COLORS["text"], padding=6)
        style.configure("TCheckbutton", background=COLORS["bg"], foreground=COLORS["text"])
        style.configure("TRadiobutton", background=COLORS["bg"], foreground=COLORS["text"])

    def set_app_icon(self):
        """Load the bundled application icon on Windows and other systems."""
        ico_path = APP_DIR_PATH / "screenshot_translator.ico"
        png_path = APP_DIR_PATH / "screenshot_translator.png"
        try:
            if os.name == "nt" and ico_path.exists():
                self.iconbitmap(str(ico_path))
            elif png_path.exists():
                self._icon_image = tk.PhotoImage(file=str(png_path))
                self.iconphoto(True, self._icon_image)
        except tk.TclError:
            pass

    def build_ui(self):
        self.build_menu()
        self.columnconfigure(0, weight=1)
        self.rowconfigure(1, weight=1)
        header = ttk.Frame(self, padding=(18, 15, 18, 8))
        header.grid(row=0, column=0, sticky="ew")
        ttk.Label(header, text="▣  Screenshot Translation", style="Header.TLabel").pack(anchor="w")
        ttk.Label(header, text="AI Vision  •  English text + Russian translation in Word",
                  style="Subtitle.TLabel").pack(anchor="w", pady=(3, 0))

        nb = ttk.Notebook(self)
        nb.grid(row=1, column=0, sticky="nsew", padx=14, pady=5)
        main = ttk.Frame(nb, padding=16)
        settings = ttk.Frame(nb, padding=16)
        nb.add(main, text="Translation")
        nb.add(settings, text="API and Prompt Settings")
        self.build_main(main)
        self.build_settings(settings)

        self.status = tk.StringVar(value="Ready. Select a capture region first.")
        ttk.Label(self, textvariable=self.status, relief="sunken", anchor="w", padding=6).grid(
            row=2, column=0, sticky="ew", padx=14, pady=(0, 12))

    def build_menu(self):
        menu = tk.Menu(self)
        help_menu = tk.Menu(menu, tearoff=False)
        help_menu.add_command(label="About", command=self.show_about)
        help_menu.add_command(label="How to Use", command=self.show_help)
        menu.add_cascade(label="Help", menu=help_menu)
        self.config(menu=menu)

    def show_about(self):
        self.show_help_window("about")

    def show_help(self):
        self.show_help_window("help")

    def show_help_window(self, mode="help"):
        window = tk.Toplevel(self)
        window.title("About — Screenshot Translator" if mode == "about" else "How to Use — Screenshot Translator")
        window.geometry("760x650")
        window.minsize(680, 560)
        window.configure(bg=COLORS["bg"])
        window.transient(self)
        window.grab_set()
        try:
            if os.name == "nt":
                window.iconbitmap(str(APP_DIR_PATH / "screenshot_translator.ico"))
        except tk.TclError:
            pass

        header = tk.Frame(window, bg=COLORS["primary"], height=94)
        header.pack(fill="x")
        header_title = "▣  Screenshot Translator — About" if mode == "about" else "▣  Screenshot Translator — How to Use"
        tk.Label(header, text=header_title, bg=COLORS["primary"], fg="white",
                 font=("Segoe UI", 22, "bold")).pack(anchor="w", padx=28, pady=(18, 0))
        header_subtitle = "Description, author, and contacts  •  version 1.0" if mode == "about" else "Step-by-step application instructions"
        tk.Label(header, text=header_subtitle, bg=COLORS["primary"], fg="#dbeafe",
                 font=("Segoe UI", 10)).pack(anchor="w", padx=30, pady=(3, 15))

        body = tk.Frame(window, bg=COLORS["bg"])
        body.pack(fill="both", expand=True, padx=24, pady=18)
        if mode == "about":
            self.help_card(body, "ABOUT THE APPLICATION", "Screenshot Translator captures the selected screen area, sends the image to an AI vision model, and saves the recognized English sentence together with its Russian translation in a Word document.\n\nVersion: 1.0", "◆", 0)
            self.help_card(body, "FEATURES", "• Select a screen area once and capture it repeatedly\n• Use one cumulative Word document or create a separate document\n• Configure the prompt, model, API, and Word margins\n• Light theme and built-in help", "✓", 1)
            self.help_card(body, "AUTHOR", "👨‍💻 Developer: Sergey Chekryzhov\n📧 Email: sergeyh510@gmail.com\n🐙 GitHub: https://github.com/sergeyh510-alt\n💡 Created to make screenshot translation convenient for users", "♥", 2)
        else:
            self.help_card(body, "STEP 1 — SETTINGS", "Open the “API and Prompt Settings” tab, enter the API URL, key, vision model, and prompt, then click “Save Settings.”", "1", 0)
            self.help_card(body, "STEP 2 — REGION", "Click “Select Region with Mouse” and drag a rectangle around the text once. The coordinates will be saved.", "2", 1)
            self.help_card(body, "STEP 3 — TRANSLATION", "Click “Capture Screenshot and Translate.” The application window is temporarily hidden so it is not included in the screenshot, then it returns after the capture.", "3", 2)
            self.help_card(body, "STEP 4 — WORD", "Choose a cumulative or separate document, configure the Word margins, and clear the prompt-to-open-file checkbox if you want to save results without additional questions.", "4", 3)
            self.help_card(body, "SETTINGS LOCATION", "%USERPROFILE%\\.screenshot_translator\\settings.xml\n\nThe API key is stored locally and is sent only to the API address you specify.", "⚙", 4)
        tk.Label(window, text="Made for convenient translation of text from screenshots", bg=COLORS["bg"], fg=COLORS["muted"],
                 font=("Segoe UI", 9)).pack(pady=(0, 14))
        ttk.Button(window, text="Close", command=window.destroy).pack(pady=(0, 18))

    @staticmethod
    def help_card(parent, title, text, mark, row):
        card = tk.Frame(parent, bg=COLORS["surface"], highlightbackground=COLORS["border"], highlightthickness=1)
        card.grid(row=row, column=0, sticky="ew", pady=(0, 10))
        parent.columnconfigure(0, weight=1)
        tk.Label(card, text=mark, bg=COLORS["surface"], fg=COLORS["primary"], font=("Segoe UI", 18, "bold"),
                 width=3).pack(side="left", anchor="n", padx=(12, 2), pady=12)
        content = tk.Frame(card, bg=COLORS["surface"])
        content.pack(side="left", fill="both", expand=True, padx=(0, 15), pady=10)
        tk.Label(content, text=title, bg=COLORS["surface"], fg=COLORS["primary_dark"],
                 font=("Segoe UI", 10, "bold"), anchor="w").pack(fill="x")
        tk.Label(content, text=text, bg=COLORS["surface"], fg=COLORS["text"], justify="left", anchor="w",
                 wraplength=650, font=("Segoe UI", 10)).pack(fill="x", pady=(4, 0))

    def build_main(self, parent):
        parent.columnconfigure(0, weight=1)
        ttk.Label(parent, text="1. Capture Region", font=("Arial", 12, "bold")).grid(row=0, column=0, sticky="w")
        controls = ttk.Frame(parent)
        controls.grid(row=1, column=0, sticky="ew", pady=(8, 18))
        self.region_label = ttk.Label(controls, text="No region selected")
        self.region_label.pack(side="left")
        ttk.Button(controls, text="Select Region with Mouse", command=self.select_region).pack(side="right")

        ttk.Separator(parent).grid(row=2, column=0, sticky="ew", pady=(0, 18))
        ttk.Label(parent, text="2. Output", font=("Arial", 12, "bold")).grid(row=3, column=0, sticky="w")
        mode = ttk.Frame(parent)
        mode.grid(row=4, column=0, sticky="ew", pady=8)
        self.mode = tk.StringVar(value=self.settings.get("mode", "append"))
        ttk.Radiobutton(mode, text="Append to one cumulative Word document", variable=self.mode, value="append").pack(anchor="w")
        ttk.Radiobutton(mode, text="Create a separate Word document for each screenshot", variable=self.mode, value="separate").pack(anchor="w")
        ttk.Label(parent, text="Word file for cumulative mode:").grid(row=5, column=0, sticky="w", pady=(12, 3))
        out = ttk.Frame(parent)
        out.grid(row=6, column=0, sticky="ew")
        out.columnconfigure(0, weight=1)
        self.output = tk.StringVar(value=self.settings.get("output_docx", DEFAULT_SETTINGS["output_docx"]))
        ttk.Entry(out, textvariable=self.output).grid(row=0, column=0, sticky="ew")
        ttk.Button(out, text="Choose…", command=self.choose_output).grid(row=0, column=1, padx=(7, 0))
        self.translate_btn = ttk.Button(parent, text="▣  Capture Screenshot and Translate", style="Accent.TButton", command=self.translate)
        self.translate_btn.grid(row=7, column=0, sticky="ew", pady=(22, 8), ipady=8)
        ttk.Checkbutton(parent, text="Ask whether to open the Word document after saving",
                        variable=self.ask_open_document).grid(row=8, column=0, sticky="w", pady=(2, 8))
        self.preview = ttk.Label(parent, text="The result will appear here after translation.", foreground="#555")
        self.preview.grid(row=9, column=0, sticky="nw", pady=10)

    def build_settings(self, parent):
        parent.columnconfigure(1, weight=1)
        fields = [("API URL:", self.api_url_var, ""),
                  ("API key:", self.api_key_var, "•"),
                  ("Model name:", self.model_var, "")]
        for i, (label, variable, mask) in enumerate(fields):
            ttk.Label(parent, text=label).grid(row=i, column=0, sticky="nw", pady=7, padx=(0, 12))
            entry = ttk.Entry(parent, textvariable=variable, show=mask)
            entry.grid(row=i, column=1, sticky="ew", pady=7)
        ttk.Label(parent, text="Model prompt:").grid(row=3, column=0, sticky="nw", pady=7, padx=(0, 12))
        self.prompt = tk.Text(parent, height=9, wrap="word")
        self.prompt.grid(row=3, column=1, sticky="nsew", pady=7)
        self.prompt.insert("1.0", self.settings.get("prompt", DEFAULT_SETTINGS["prompt"]))
        parent.rowconfigure(3, weight=1)
        ttk.Label(parent, text="Word margins, cm:").grid(row=4, column=0, sticky="nw", pady=7, padx=(0, 12))
        margins = ttk.Frame(parent)
        margins.grid(row=4, column=1, sticky="w", pady=7)
        for label, variable in (("Top", self.margin_top_var), ("Bottom", self.margin_bottom_var),
                                ("Left", self.margin_left_var), ("Right", self.margin_right_var)):
            ttk.Label(margins, text=label).pack(side="left", padx=(0, 3))
            ttk.Entry(margins, textvariable=variable, width=6).pack(side="left", padx=(0, 12))
        ttk.Label(parent, text="AI uses an OpenAI-compatible (RouterAI) POST /chat/completions endpoint. A vision model is required.",
                  foreground="#555", wraplength=570).grid(row=5, column=1, sticky="w", pady=(4, 10))
        ttk.Button(parent, text="Save Settings", command=self.persist_settings).grid(row=6, column=1, sticky="e")

    def persist_settings(self):
        self.settings.update({"api_url": self.api_url_var.get().strip(), "api_key": self.api_key_var.get().strip(),
                              "model": self.model_var.get().strip(), "prompt": self.prompt.get("1.0", "end").strip(),
                              "output_docx": self.output.get().strip(), "mode": self.mode.get(),
                              "margin_top": self.margin_top_var.get().strip(),
                              "margin_bottom": self.margin_bottom_var.get().strip(),
                              "margin_left": self.margin_left_var.get().strip(),
                              "margin_right": self.margin_right_var.get().strip(),
                              "ask_open_document": self.ask_open_document.get()})
        save_settings(self.settings)
        self.status.set("Settings saved locally.")

    def select_region(self):
        self.withdraw()
        self.after(250, lambda: RegionSelector(self, self.region_selected))

    def region_selected(self, region):
        self.deiconify()
        self.region = region
        self.region_label.configure(text=f"Saved: x={region[0]}, y={region[1]}, width={region[2]-region[0]}, height={region[3]-region[1]}")
        self.status.set("Region saved. You can click the translation button at any time.")

    def choose_output(self):
        path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word document", "*.docx")],
                                            initialfile="ScreenshotTranslations.docx")
        if path:
            self.output.set(path)

    def translate(self):
        if self.busy:
            return
        if not self.region:
            messagebox.showwarning("No region selected", "First click “Select Region with Mouse” and select the text.")
            return
        self.persist_settings()
        if not self.settings.get("api_key") or not self.settings.get("model"):
            messagebox.showwarning("Incomplete settings", "Enter the API key and the vision model name on the settings tab.")
            return
        self.busy = True
        self.translate_btn.configure(state="disabled")
        self.status.set("Capturing the region and sending the image to RouterAI…")
        # Hide the application so it cannot appear in the next screenshot.
        self.withdraw()
        self.after(250, self.start_worker_after_hide)

    def start_worker_after_hide(self):
        threading.Thread(target=self.worker, daemon=True).start()

    def worker(self):
        try:
            image = ImageGrab.grab(bbox=self.region, all_screens=True)
            # The screenshot is complete; show the app again while the API works.
            self.after(0, self.deiconify)
            buf = io.BytesIO()
            image.save(buf, format="PNG")
            encoded = base64.b64encode(buf.getvalue()).decode("ascii")
            payload = {"model": self.settings["model"], "messages": [
                {"role": "system", "content": self.settings["prompt"]},
                {"role": "user", "content": [{"type": "text", "text": "Recognize and translate the text in this image."},
                                             {"type": "image_url", "image_url": {"url": "data:image/png;base64," + encoded}}]}
            ], "temperature": 0.1}
            response = requests.post(self.settings["api_url"], headers={"Authorization": "Bearer " + self.settings["api_key"],
                                                                        "Content-Type": "application/json"},
                                     json=payload, timeout=120)
            response.raise_for_status()
            data = response.json()
            text = data["choices"][0]["message"]["content"]
            self.save_result(text, image)
            self.after(0, self.done, text)
        except Exception as exc:
            # Exception variables are cleared when the except block ends.
            # Copy the message before scheduling the callback in Tkinter.
            error_text = str(exc)
            self.after(0, self.deiconify)
            self.after(0, self.failed, error_text)

    def save_result(self, text, image):
        if self.mode.get() == "separate":
            base = Path(self.output.get()).with_suffix("")
            path = base.parent / f"{base.name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            doc = Document()
        else:
            path = Path(self.output.get())
            path.parent.mkdir(parents=True, exist_ok=True)
            if path.exists():
                try:
                    doc = Document(str(path))
                except (PackageNotFoundError, ValueError, OSError):
                    # A file with a .docx extension may be empty or may not be a real
                    # Office Open XML package. Start a clean document instead.
                    doc = Document()
            else:
                doc = Document()
        self.apply_margins(doc)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Screenshot Translation")
        run.bold = True
        run.font.size = Pt(11)
        doc.add_picture(io.BytesIO(self.image_bytes(image)), width=Inches(5.8))
        for line in [x.strip() for x in text.strip().splitlines() if x.strip()]:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(3)
        doc.add_paragraph("─" * 70)
        doc.save(path)
        self.last_path = path

    def apply_margins(self, doc):
        """Apply page margins in centimeters to every section of the document."""
        try:
            values = [float(self.margin_top_var.get().replace(",", ".")),
                      float(self.margin_bottom_var.get().replace(",", ".")),
                      float(self.margin_left_var.get().replace(",", ".")),
                      float(self.margin_right_var.get().replace(",", "."))]
            if any(value < 0 for value in values):
                raise ValueError
        except ValueError:
            raise ValueError("Word margins must be non-negative numbers in centimeters.")
        top, bottom, left, right = values
        for section in doc.sections:
            section.top_margin = Cm(top)
            section.bottom_margin = Cm(bottom)
            section.left_margin = Cm(left)
            section.right_margin = Cm(right)

    @staticmethod
    def image_bytes(image):
        b = io.BytesIO(); image.save(b, format="PNG"); return b.getvalue()

    def done(self, text):
        self.busy = False; self.translate_btn.configure(state="normal")
        self.preview.configure(text=text[:500])
        self.status.set(f"Done. Document saved: {self.last_path}")
        if self.ask_open_document.get() and messagebox.askyesno("Translation Complete", f"Open the document?\n{self.last_path}"):
            if os.name == "nt":
                os.startfile(str(self.last_path))
            elif os.name == "posix":
                subprocess.Popen(["xdg-open", str(self.last_path)])

    def failed(self, error):
        self.busy = False; self.translate_btn.configure(state="normal")
        self.status.set("Translation error.")
        messagebox.showerror("Error", error)


if __name__ == "__main__":
    App().mainloop()
